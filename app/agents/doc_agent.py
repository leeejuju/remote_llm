from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from app.agents.base_agent import BaseAgent
from app.schema.schema import QuestionAnswerState


class DocAgent(BaseAgent):
    def __init__(self, output_path="分析结果.docx"):
        self.output_path = output_path

    def write_analysis(self, analysis_result: QuestionAnswerState, title: str = "分析结果"):
        doc = Document()

        # 标题
        doc.add_heading(title, 0)

        # 添加正文
        para = doc.add_paragraph(analysis_result.answer)

        # 设置字体为微软雅黑，字号 12pt
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if getattr(analysis_result, "segmented_image_path", None):
                img_path = analysis_result.segmented_image_path
                doc.add_picture(img_path, width=Inches(4))

        # 保存 Word
        doc.save(self.output_path)
